import os
import re
import json
import base64
import urllib.request
import urllib.parse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_table_pagination(table):
    # Repeat header on new pages and prevent rows from splitting awkwardly
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def fetch_mermaid_image(mermaid_code, output_path):
    try:
        clean_code = mermaid_code.strip()
        clean_code = clean_code.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
        payload = {
            "code": clean_code,
            "mermaid": {
                "theme": "default",
                "themeVariables": {
                    "fontSize": "16px",
                    "fontFamily": "Segoe UI, Arial, sans-serif"
                }
            }
        }
        json_bytes = json.dumps(payload).encode("utf-8")
        base64_str = base64.b64encode(json_bytes).decode("ascii")
        url = f"https://mermaid.ink/img/{base64_str}"
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=30) as resp:
            data = resp.read()
            with open(output_path, "wb") as f:
                f.write(data)
            print(f"Successfully generated diagram image: {output_path} ({len(data)} bytes)")
            return True
    except Exception as e:
        print(f"Error fetching mermaid image: {e}")
        return False

def clean_text_formatting(text):
    if not text:
        return ""
    text = text.replace(r"$\rightarrow$", "→")
    text = text.replace(r"\rightarrow", "→")
    text = text.replace("&nbsp;", " ")
    text = text.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
    return text

def build_word_document():
    md_path = "docs/shimano-platform-handbook.md"
    os.makedirs("docs/images", exist_ok=True)

    if not os.path.exists(md_path):
        print(f"Markdown file {md_path} not found.")
        return
        
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()
    
    # Page setup - 0.65 inch margins (7.2 in printable width, 9.7 in printable height)
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        
    NAVY_HEX = "003366"
    BLUE_HEX = "005A9C"
    DARK_TEXT = RGBColor(35, 35, 35)
    NAVY_COLOR = RGBColor(0, 51, 102)
    BLUE_COLOR = RGBColor(0, 90, 156)
    CODE_COLOR = RGBColor(170, 30, 30)
    
    # Base tight, elegant typography
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = DARK_TEXT
    normal_style.paragraph_format.line_spacing = 1.10
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(2.5)

    # Document Header Title Block - Compact & Professional
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Shimano Experience Platform")
    title_run.font.size = Pt(21)
    title_run.font.bold = True
    title_run.font.color.rgb = NAVY_COLOR
    title_p.paragraph_format.space_after = Pt(1)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Multi-Repository Delivery Playbook & Architecture Guide")
    sub_run.font.size = Pt(12.5)
    sub_run.font.color.rgb = BLUE_COLOR
    sub_run.font.bold = True
    sub_p.paragraph_format.space_after = Pt(6)

    # Metadata callout box - Compact
    meta_tbl = doc.add_table(rows=1, cols=1)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = meta_tbl.cell(0, 0)
    cell.width = Inches(7.2)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
    
    meta_p = cell.paragraphs[0]
    meta_p.paragraph_format.space_after = Pt(1)
    meta_p.paragraph_format.line_spacing = 1.05
    
    r1 = meta_p.add_run("Target Audience: ")
    r1.bold = True
    r1.font.size = Pt(9)
    r1_t = meta_p.add_run("Development Teams, Tech Leads, Solution Architects, Security Engineers, & Leadership\n")
    r1_t.font.size = Pt(9)
    
    r2 = meta_p.add_run("Platform: ")
    r2.bold = True
    r2.font.size = Pt(9)
    r2_t = meta_p.add_run("Adobe Experience Manager as a Cloud Service (AEMaaCS), GitHub Actions, Git Subtree\n")
    r2_t.font.size = Pt(9)
    
    r3 = meta_p.add_run("Confluence Space: ")
    r3.bold = True
    r3.font.size = Pt(9)
    r3_t = meta_p.add_run("Engineering > Architecture & Delivery > Platform Playbook")
    r3_t.font.size = Pt(9)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="20" w:space="0" w:color="{NAVY_HEX}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    lines = md_text.splitlines()
    in_code_block = False
    code_lang = ""
    code_lines = []
    in_table = False
    table_lines = []
    in_callout = False
    callout_type = ""
    callout_lines = []
    diagram_count = 0

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        rows_data = []
        for line in tbl_lines:
            if re.match(r'^\s*\|?\s*:?---', line):
                continue
            cells = [clean_text_formatting(c.strip()) for c in line.strip().strip('|').split('|')]
            if cells:
                rows_data.append(cells)
        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, "CBD5E1")
        format_table_pagination(tbl)

        for row_idx, row in enumerate(rows_data):
            for col_idx in range(num_cols):
                c_text = row[col_idx] if col_idx < len(row) else ""
                t_cell = tbl.cell(row_idx, col_idx)
                set_cell_margins(t_cell, top=60, bottom=60, left=100, right=100)
                t_p = t_cell.paragraphs[0]
                t_p.paragraph_format.space_before = Pt(0)
                t_p.paragraph_format.space_after = Pt(1.5)
                t_p.paragraph_format.line_spacing = 1.05

                if row_idx == 0:
                    set_cell_background(t_cell, NAVY_HEX)
                    run = t_p.add_run(re.sub(r'[\*`]', '', c_text))
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(t_cell, "F8FAFC")
                    else:
                        set_cell_background(t_cell, "FFFFFF")
                    
                    render_inline_markdown(t_p, c_text, font_size=Pt(8.5))
        
        doc.add_paragraph().paragraph_format.space_after = Pt(3)

    def render_inline_markdown(para, text, font_size=Pt(10)):
        text = clean_text_formatting(text)
        text = re.sub(r'<br\s*/?>', '\n', text)
        tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith('**') and tok.endswith('**'):
                r = para.add_run(tok[2:-2])
                r.bold = True
                r.font.size = font_size
            elif tok.startswith('*') and tok.endswith('*'):
                r = para.add_run(tok[1:-1])
                r.italic = True
                r.font.size = font_size
            elif tok.startswith('`') and tok.endswith('`'):
                r = para.add_run(tok[1:-1])
                r.font.name = 'Consolas'
                r.font.size = Pt(font_size.pt * 0.90)
                r.font.color.rgb = CODE_COLOR
            else:
                clean_tok = re.sub(r'<[^>]+>', '', tok)
                r = para.add_run(clean_tok)
                r.font.size = font_size

    def flush_callout(c_type, c_lines):
        if not c_lines:
            return
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(7.2)
        
        bg_hex = "F0FDF4" if c_type == "TIP" else ("FEF2F2" if c_type == "CAUTION" else ("FFFBEB" if c_type == "IMPORTANT" else "F0F7FF"))
        border_hex = "16A34A" if c_type == "TIP" else ("DC2626" if c_type == "CAUTION" else ("D97706" if c_type == "IMPORTANT" else BLUE_HEX))
        
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=80, bottom=80, left=130, right=130)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(tcBorders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.05
        
        title_run = p.add_run(f"📌 {c_type}: ")
        title_run.bold = True
        title_run.font.color.rgb = NAVY_COLOR
        title_run.font.size = Pt(9.5)
        
        full_text = "\n".join(c_lines)
        render_inline_markdown(p, full_text, font_size=Pt(9))
        doc.add_paragraph().paragraph_format.space_after = Pt(3)

    def flush_code(c_lang, c_lines):
        nonlocal diagram_count
        if not c_lines:
            return
        if c_lang == "mermaid":
            diagram_count += 1
            mermaid_text = "\n".join(c_lines)
            img_path = f"docs/images/diagram_{diagram_count}.png"
            if fetch_mermaid_image(mermaid_text, img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.keep_with_next = True
                
                # Proportional image scaling to prevent accidental overflow
                if diagram_count == 1:
                    # Width 5.8 inches -> Height = 4.4 inches (fits comfortably on same page with header)
                    doc.add_picture(img_path, width=Inches(5.8))
                else:
                    # Sequence diagram -> Width 5.4 inches -> Height = 3.8 inches
                    doc.add_picture(img_path, width=Inches(5.4))
                    
                doc.add_paragraph().paragraph_format.space_after = Pt(3)
                return

        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(7.2)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/></w:tcBorders>')
        tcPr.append(tcBorders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.02
        
        code_text = "\n".join(c_lines)
        code_text = clean_text_formatting(code_text)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.0)
        run.font.color.rgb = RGBColor(30, 30, 30)
        doc.add_paragraph().paragraph_format.space_after = Pt(3)

    lines_to_process = lines[7:]
    
    for line in lines_to_process:
        if line.startswith("```"):
            if in_code_block:
                flush_code(code_lang, code_lines)
                in_code_block = False
                code_lang = ""
                code_lines = []
            else:
                if in_table:
                    flush_table(table_lines)
                    in_table = False
                    table_lines = []
                if in_callout:
                    flush_callout(callout_type, callout_lines)
                    in_callout = False
                    callout_lines = []
                in_code_block = True
                code_lang = line[3:].strip()
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        callout_match = re.match(r'^>\s*\[!(NOTE|TIP|IMPORTANT|WARNING|CAUTION)\]', line)
        if callout_match:
            if in_table:
                flush_table(table_lines)
                in_table = False
                table_lines = []
            in_callout = True
            callout_type = callout_match.group(1)
            callout_lines = []
            continue

        if in_callout:
            if line.startswith(">"):
                callout_lines.append(line.lstrip("> ").strip())
                continue
            else:
                flush_callout(callout_type, callout_lines)
                in_callout = False
                callout_lines = []

        if re.match(r'^\s*\|.*\|\s*$', line):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            continue
        else:
            if in_table:
                flush_table(table_lines)
                in_table = False
                table_lines = []

        if line.startswith("## "):
            h_text = clean_text_formatting(line[3:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = NAVY_COLOR
            continue

        if line.startswith("### "):
            h_text = clean_text_formatting(line[4:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.size = Pt(11.5)
            run.bold = True
            run.font.color.rgb = BLUE_COLOR
            continue

        if line.startswith("#### "):
            h_text = clean_text_formatting(line[5:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.size = Pt(10.5)
            run.bold = True
            run.font.color.rgb = DARK_TEXT
            continue

        if line.strip() == "---":
            continue

        list_match = re.match(r'^\s*([0-9]+\.|\-|\*)\s+(.*)$', line)
        if list_match:
            bullet_char = list_match.group(1)
            item_text = list_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.left_indent = Inches(0.2)
            
            b_run = p.add_run(f"{bullet_char} ")
            b_run.bold = True
            b_run.font.color.rgb = NAVY_COLOR
            render_inline_markdown(p, item_text)
            continue

        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2.5)
            render_inline_markdown(p, line.strip())

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lang, code_lines)
    if in_callout:
        flush_callout(callout_type, callout_lines)

    saved = False
    for candidate in [
        "docs/Shimano-Experience-Platform-Handbook.docx",
        "docs/Shimano-Experience-Platform-Handbook-v2.docx",
        "docs/Shimano-Experience-Platform-Handbook-v3.docx",
        "docs/Shimano-Experience-Platform-Handbook-Compact.docx"
    ]:
        try:
            doc.save(candidate)
            print(f"SUCCESS: Saved tight presentation Word document to {candidate} ({os.path.getsize(candidate)} bytes)")
            saved = True
            break
        except PermissionError:
            continue

    if not saved:
        fallback = "docs/Shimano-Experience-Platform-Handbook-Final.docx"
        doc.save(fallback)
        print(f"SUCCESS: Saved tight presentation Word document to {fallback} ({os.path.getsize(fallback)} bytes)")

if __name__ == "__main__":
    build_word_document()
